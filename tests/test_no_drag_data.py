"""An .ork with NO stored drag data must flow through EVERY page's backend
(Simulate, Monte Carlo, RCSM cases, Analysis, Compare Designs, report,
LASC zip) using a constant placeholder Cd - never crash, always labelled
LOW CONFIDENCE. A tester hit "power_off_drag/power_on_drag are required"
on Monte Carlo / RCSM / Analysis / Export because only Simulate knew how
to fall back; the fix is LoadResult.drag_curves() as the single source.
"""
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from bup_rocketpy import analysis, compare_designs, lasc_package, monte_carlo, rcsm_cases, report, translate
from bup_rocketpy.gui import pipeline

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_ORK = os.path.join(REPO_ROOT, "reference", "prometeo_mission44", "data", "ork", "PrometeoLasc2026.ork")
ENG_PATH = os.path.join(REPO_ROOT, "reference", "prometeo_mission44", "data", "motors", "Icarus_I_K519.eng")
OUT_DIR = os.path.join(REPO_ROOT, "outputs", "test_no_drag_data")
DRY_MASS_KG, DRY_CG_M = 5.6622, 0.6279


def _make_no_drag_ork():
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(SRC_ORK, encoding="utf-8") as f:
        data = f.read()
    stripped = re.sub(r"<databranch\b.*?</databranch>", "", data, flags=re.DOTALL)
    assert stripped != data, "fixture must actually contain a databranch to strip"
    path = os.path.join(OUT_DIR, "NoDragData.ork")
    with open(path, "w", encoding="utf-8") as f:
        f.write(stripped)
    return path


def _load():
    ork = _make_no_drag_ork()
    return ork, pipeline.load_files(ork, ENG_PATH, outputs_dir=OUT_DIR)


def test_load_marks_placeholder_and_writes_real_csv_files():
    _, lr = _load()
    assert lr.drag_is_placeholder is True
    assert "LOW CONFIDENCE" in lr.drag_curve_source
    off, on = lr.drag_curves()
    assert os.path.exists(off) and os.path.exists(on)
    assert pipeline.LOW_CONFIDENCE_DRAG_BANNER.startswith("No drag (Cd) data in this file")


def test_every_backend_completes_without_drag_data():
    ork, lr = _load()
    off, on = lr.drag_curves()

    sim = pipeline.run_simulation(lr, OUT_DIR, dry_mass_override_kg=DRY_MASS_KG, dry_cg_override_m=DRY_CG_M)
    assert sim.apogee_agl_m > 100

    mass_est = translate.MassEstimate(DRY_MASS_KG, DRY_CG_M, "test")
    i_ax, i_tr = translate.estimate_dry_inertia(lr.parsed_ork, mass_est)
    radius = next(t.radius for t in lr.parsed_ork.body_tubes if t.radius)

    uncertainties = monte_carlo.default_uncertainties(DRY_MASS_KG, 1855.9, lr.parsed_ork.launch.wind_average_ms)
    mc = monte_carlo.run_monte_carlo(
        lr.parsed_ork, lr.parsed_eng, lr.eng_path, *lr.drag_curves(), DRY_MASS_KG, DRY_CG_M, i_ax, i_tr, radius,
        uncertainties, n_simulations=8, output_dir=OUT_DIR, max_workers=4, seed=1,
    )
    assert mc.n_completed == 8

    cases = rcsm_cases.run_all_cases(lr.parsed_ork, lr.parsed_eng, lr.eng_path, *lr.drag_curves(), DRY_MASS_KG, DRY_CG_M)
    assert "Ballistic" in cases and "Nominal" in cases

    sweep = analysis.weathercocking_sweep(lr.parsed_ork, lr.parsed_eng, lr.eng_path, *lr.drag_curves(), DRY_MASS_KG, DRY_CG_M, i_ax, i_tr, radius)
    assert sweep is not None

    cmp_result = compare_designs.compare_designs(ork, ENG_PATH, ork, ENG_PATH, outputs_dir=OUT_DIR)
    assert cmp_result.design_a.apogee_agl_m > 100

    data = report.build_report_data("44", "Test", lr, sim, cases, mc, uncertainties, OUT_DIR)
    docx_path = report.generate_docx(os.path.join(OUT_DIR, "report.docx"), data)
    from docx import Document
    doc_text = "\n".join(p.text for p in Document(docx_path).paragraphs) + "\n".join(
        c.text for t in Document(docx_path).tables for r in t.rows for c in r.cells
    )
    assert "LOW CONFIDENCE" in doc_text
    pdf_path = report.generate_pdf(os.path.join(OUT_DIR, "report.pdf"), data)
    assert os.path.getsize(pdf_path) > 1000

    zip_path = os.path.join(OUT_DIR, "no_drag.zip")
    lasc_package.build_lasc_zip(
        zip_path, "44", lr.parsed_ork, lr.parsed_eng, lr.eng_path, ork, *lr.drag_curves(),
        DRY_MASS_KG, DRY_CG_M, i_ax, i_tr, radius,
        cases=[("Ballistic", False), ("Nominal", True)],
        eng_filename="Icarus_I_K519.eng", ork_filename="NoDragData.ork",
        drag_is_placeholder=lr.drag_is_placeholder,
    )
    with zipfile.ZipFile(zip_path) as zf:
        names = zf.namelist()
        assert any(n.endswith("power_off_drag.csv") for n in names)
        readme = next(zf.read(n).decode("utf-8") for n in names if n.lower().endswith("readme.md"))
        script = next(zf.read(n).decode("utf-8") for n in names if "Nominal" in n and n.endswith(".py"))
    assert "LOW CONFIDENCE" in readme
    assert "LOW CONFIDENCE" in script
